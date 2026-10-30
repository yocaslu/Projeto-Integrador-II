"""Leitura da planilha do Anuário (dados/brutos/anuario-2026.xlsx).

Segue docs/importacao.md, seções 4.2 e 4.3: localiza a âncora, expande as células mescladas,
resolve cada coluna pelo caminho de cabeçalho, lê as linhas até "Fonte:" (pulando as em branco)
e devolve `Linha`s sem gravar nada no banco. Variações e proporções publicadas, e as tabelas de
conferência (T01, T09), saem como `ValorPublicado`, que só alimenta o relatório de validação.
"""

import re
from collections.abc import Mapping
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.workbook import Workbook

from apps.importacao.erros import ErroImportacao
from apps.importacao.linhas import Chave, Linha, Localidade
from apps.importacao.mapeamento import GRUPOS, GRUPOS_GRAVADOS, Bloco, Tabela
from apps.importacao.normalizacao import (
    Situacao,
    chave_de_nome,
    normalizar_celula,
    normalizar_taxa,
    resolver_localidade,
)
from apps.registros.escolhas import Unidade

LINHAS_PROCURA_ANCORA = 30
MARCA = re.compile(r"\s*\((\d+)\)")
NOTA = re.compile(r"^\((\d+)\)\s*(.+)$", re.DOTALL)
FONTE = re.compile(r"^\s*Fonte:")
ANO = re.compile(r"^(\d{4})\b")


def abrir(caminho: Path) -> Workbook:
    """Abre a planilha só com valores. Sem read_only: as células mescladas são necessárias."""
    return openpyxl.load_workbook(caminho, data_only=True)


class Grade:
    """Valores de uma aba, com o valor de cada célula mesclada replicado no intervalo inteiro."""

    def __init__(self, ws):
        self.ws = ws
        self.max_row = ws.max_row
        self.max_col = ws.max_column
        self._mescladas: dict[tuple[int, int], object] = {}
        for intervalo in ws.merged_cells.ranges:
            valor = ws.cell(intervalo.min_row, intervalo.min_col).value
            for linha in range(intervalo.min_row, intervalo.max_row + 1):
                for coluna in range(intervalo.min_col, intervalo.max_col + 1):
                    self._mescladas[(linha, coluna)] = valor

    def valor(self, linha: int, coluna: int) -> object:
        if (linha, coluna) in self._mescladas:
            return self._mescladas[(linha, coluna)]
        return self.ws.cell(linha, coluna).value


@dataclass(frozen=True)
class Coluna:
    bloco: Bloco
    grupo: str
    ano: int
    marcas: tuple[str, ...]  # marcas de nota que valem para o dado (não as da base da taxa)


@dataclass(frozen=True)
class ValorPublicado:
    """Valor lido só para o relatório: variação, proporção ou tabela de conferência."""

    tabela: str
    origem: str  # célula, ex.: "T03!F10"
    grupo: str  # valor | taxa | variacao | proporcao
    chave: Chave
    valor: Decimal | None
    situacao: Situacao = Situacao.INFORMADO  # "-" e "..." são marcadores, não números


@dataclass
class TabelaLida:
    tabela: Tabela
    linhas: list[Linha] = field(default_factory=list)
    publicados: list[ValorPublicado] = field(default_factory=list)
    vazias: list[str] = field(default_factory=list)  # células vazias na área de dados
    fonte: str = ""  # linha "Fonte:" completa


def coordenada(linha: int, coluna: int) -> str:
    return f"{get_column_letter(coluna)}{linha}"


def texto(valor: object) -> str:
    """Texto de uma célula de cabeçalho, com espaços e quebras de linha normalizados."""
    return re.sub(r"\s+", " ", str(valor)).strip() if valor is not None else ""


def separar_marcas(rotulo: object) -> tuple[str, list[str]]:
    """'Rio Grande do Norte (3) (4)' → ('Rio Grande do Norte', ['3', '4'])."""
    bruto = texto(rotulo)
    return MARCA.sub("", bruto).strip(), MARCA.findall(bruto)


def achar_ancora(grade: Grade, regex: str, aba: str, coluna: int = 1) -> int:
    """1ª linha cuja coluna de rótulos (A, ou B na edição de 2024) casa a regex inteira (^…$)."""
    for linha in range(1, min(grade.max_row, LINHAS_PROCURA_ANCORA) + 1):
        valor = grade.ws.cell(linha, coluna).value
        if valor is not None and re.match(regex, str(valor)):
            return linha
    raise ErroImportacao(f"{aba}: âncora não encontrada ({regex}).")


def achar_fonte(grade: Grade, inicio: int, aba: str) -> int:
    for linha in range(inicio, grade.max_row + 1):
        valor = grade.ws.cell(linha, 1).value
        if valor is not None and FONTE.match(str(valor)):
            return linha
    raise ErroImportacao(f"{aba}: linha 'Fonte:' não encontrada.")


def ler_notas(grade: Grade, linha_fonte: int) -> dict[str, str]:
    """Notas numeradas do rodapé: {"3": "Atualização das informações ..."}."""
    notas = {}
    for linha in range(linha_fonte + 1, grade.max_row + 1):
        valor = grade.ws.cell(linha, 1).value
        if valor is None:
            continue
        nota = NOTA.match(str(valor).strip())
        if nota:
            notas[nota.group(1)] = re.sub(r"\s+", " ", nota.group(2)).strip()
    return notas


def caminhos(
    grade: Grade, linha_ancora: int, niveis: int, coluna_rotulo: int = 1
) -> dict[int, tuple[str, ...]]:
    """Caminho de cabeçalho de cada coluna com cabeçalho (depois da coluna de rótulos)."""
    resultado = {}
    for coluna in range(coluna_rotulo + 1, grade.max_col + 1):
        caminho = tuple(
            texto(grade.valor(linha, coluna))
            for linha in range(linha_ancora, linha_ancora + niveis)
        )
        if any(caminho):
            resultado[coluna] = caminho
    return resultado


def _casa_caminho(regexes: tuple[str, ...], caminho: tuple[str, ...]) -> bool:
    """Cada regex precisa casar com um nível do caminho, em ordem."""
    inicio = 0
    for regex in regexes:
        for indice in range(inicio, len(caminho)):
            if re.search(regex, caminho[indice], re.IGNORECASE):
                inicio = indice + 1
                break
        else:
            return False
    return True


def _grupo(nivel: str) -> str | None:
    for regex, grupo in GRUPOS.items():
        if re.search(regex, nivel, re.IGNORECASE):
            return grupo
    return None


def _conferir_base(
    rotulo_grupo: str, marcas: list[str], notas: Mapping[str, str], base: str | None, local: str
) -> None:
    """A base da taxa vem do rótulo do grupo ou da nota marcada nele (docs/importacao.md, 4.3)."""
    if base is None:
        raise ErroImportacao(f"{local}: coluna de taxa em indicador sem base de taxa.")
    if marcas:
        declarada = notas.get(marcas[0])
        if declarada is None:
            raise ErroImportacao(f"{local}: nota ({marcas[0]}) do grupo de taxa não encontrada.")
    else:
        declarada = rotulo_grupo
    if chave_de_nome(base) not in chave_de_nome(declarada):
        raise ErroImportacao(
            f"{local}: a base da taxa declarada ({declarada!r}) não é a do indicador ({base})."
        )


def resolver_colunas(
    grade: Grade,
    linha_ancora: int,
    tabela: Tabela,
    notas: Mapping[str, str],
    base_taxa: Mapping[str, str | None],
) -> dict[int, Coluna]:
    """Bloco, grupo, ano e marcas de cada coluna. Coluna órfã ou ambígua aborta a importação."""
    colunas: dict[int, Coluna] = {}
    colunas_lidas = caminhos(grade, linha_ancora, tabela.linhas_cabecalho, tabela.coluna_rotulo)
    for coluna, caminho in colunas_lidas.items():
        local = f"{tabela.aba}!{get_column_letter(coluna)}{linha_ancora}"
        blocos = [b for b in tabela.blocos if _casa_caminho(b.caminho, caminho)]
        if not blocos:
            if any(re.search(r, nivel, re.IGNORECASE) for r in tabela.ignorar for nivel in caminho):
                continue
            raise ErroImportacao(f"{local}: coluna sem bloco no mapeamento: {caminho}.")
        if len(blocos) > 1:
            raise ErroImportacao(f"{local}: coluna casa com mais de um bloco: {caminho}.")
        bloco = blocos[0]

        rotulo_grupo, marcas_grupo = separar_marcas(caminho[-2])
        grupo = _grupo(rotulo_grupo)
        if grupo is None or grupo not in bloco.grupos:
            raise ErroImportacao(f"{local}: grupo {caminho[-2]!r} não previsto no bloco.")

        ultimo, marcas_ano = separar_marcas(caminho[-1])
        ano_lido = ANO.match(ultimo)
        if ano_lido:
            ano = int(ano_lido.group(1))
        elif grupo == "variacao":
            ano = max(tabela.anos)  # variação do último ano em relação ao anterior
        else:
            raise ErroImportacao(f"{local}: ano não encontrado no cabeçalho: {caminho}.")
        if ano not in tabela.anos:
            raise ErroImportacao(f"{local}: ano {ano} fora dos anos da tabela.")

        if grupo == "taxa":
            _conferir_base(rotulo_grupo, marcas_grupo, notas, base_taxa[bloco.indicador], local)
            marcas_grupo = []  # a nota da base não descreve o dado
        colunas[coluna] = Coluna(bloco, grupo, ano, tuple(marcas_grupo + marcas_ano))

    _conferir_colunas(colunas, tabela)
    return colunas


def _conferir_colunas(colunas: Mapping[int, Coluna], tabela: Tabela) -> None:
    """Todo bloco precisa ter, em cada ano, uma coluna de cada grupo lido (menos variação)."""
    presentes = {(c.bloco, c.grupo, c.ano) for c in colunas.values()}
    for bloco in tabela.blocos:
        for grupo in bloco.grupos - {"variacao"}:
            for ano in tabela.anos:
                if (bloco, grupo, ano) not in presentes:
                    raise ErroImportacao(
                        f"{tabela.aba}: falta a coluna {grupo} de {ano} do bloco "
                        f"{bloco.indicador} {bloco.recorte or ''}."
                    )


def juntar_notas(notas: Mapping[str, str], marcas: set[str]) -> str:
    """Textos das notas marcadas, em ordem numérica, separados por espaço."""
    return " ".join(notas[m] for m in sorted(marcas, key=int) if m in notas)


def ler_tabela(
    wb: Workbook,
    tabela: Tabela,
    unidade_padrao: Mapping[str, Unidade],
    base_taxa: Mapping[str, str | None],
) -> TabelaLida:
    """Lê uma tabela do mapeamento; qualquer inconsistência levanta ErroImportacao."""
    if tabela.aba not in wb.sheetnames:
        raise ErroImportacao(f"Aba {tabela.aba} não encontrada na planilha.")
    grade = Grade(wb[tabela.aba])
    linha_ancora = achar_ancora(grade, tabela.ancora, tabela.aba, tabela.coluna_rotulo)
    linha_fonte = achar_fonte(grade, linha_ancora + tabela.linhas_cabecalho, tabela.aba)
    notas = ler_notas(grade, linha_fonte)
    colunas = resolver_colunas(grade, linha_ancora, tabela, notas, base_taxa)

    lida = TabelaLida(tabela, fonte=texto(grade.ws.cell(linha_fonte, 1).value))
    linhas: dict[Chave, Linha] = {}
    localidades: set[Localidade] = set()

    for r in range(linha_ancora + tabela.linhas_cabecalho, linha_fonte):
        rotulo = grade.ws.cell(r, tabela.coluna_rotulo).value
        if rotulo is None or texto(rotulo) == "":
            # Na edição de 2024, a linha do Brasil tem o rótulo na coluna A, mesclada com a B.
            rotulo = grade.ws.cell(r, 1).value
        if rotulo is None or texto(rotulo) == "":
            continue  # a T02 tem linhas em branco entre as regiões
        nome, marcas_linha = separar_marcas(rotulo)
        try:
            localidade = resolver_localidade(nome)
        except KeyError as erro:
            raise ErroImportacao(f"{tabela.aba}!A{r}: {erro.args[0]}") from None
        if localidade in localidades:
            raise ErroImportacao(f"{tabela.aba}!A{r}: localidade repetida ({nome}).")
        localidades.add(localidade)

        for c, coluna in colunas.items():
            bruto = grade.ws.cell(r, c).value
            origem = f"{tabela.aba}!{coordenada(r, c)}"
            if bruto is None or (isinstance(bruto, str) and bruto.strip() == ""):
                lida.vazias.append(origem)
            try:
                if coluna.grupo == "valor":
                    celula = normalizar_celula(bruto)
                else:
                    celula = normalizar_taxa(bruto)
            except ValueError as erro:
                raise ErroImportacao(f"{origem}: {erro}") from None

            bloco = coluna.bloco
            unidade = bloco.unidade or unidade_padrao[bloco.indicador]
            chave = Chave(localidade, bloco.indicador, bloco.recorte, unidade, coluna.ano)
            if tabela.conferencia or coluna.grupo not in GRUPOS_GRAVADOS:
                lida.publicados.append(
                    ValorPublicado(
                        tabela.aba, origem, coluna.grupo, chave, celula.valor, celula.situacao
                    )
                )
                continue

            linha = linhas.setdefault(chave, Linha(chave=chave))
            if coluna.grupo == "valor":
                linha.valor, linha.situacao, linha.origem = celula.valor, celula.situacao, origem
            else:
                linha.taxa_publicada, linha.origem_taxa = celula.valor, origem
            linha.marcas |= set(marcas_linha) | set(coluna.marcas)

    for linha in linhas.values():
        linha.observacao = juntar_notas(notas, linha.marcas)
    lida.linhas = sorted(linhas.values(), key=lambda linha: linha.chave)
    _conferir_contagem(lida, localidades)
    return lida


def _conferir_contagem(lida: TabelaLida, localidades: set[Localidade]) -> None:
    """Localidades × anos × blocos gravados; nunca pular linha em silêncio (regra 7)."""
    tabela = lida.tabela
    if len(localidades) != tabela.localidades_esperadas:
        raise ErroImportacao(
            f"{tabela.aba}: {len(localidades)} localidades lidas; esperadas "
            f"{tabela.localidades_esperadas}."
        )
    if tabela.conferencia:
        return
    esperadas = len(localidades) * len(tabela.anos) * len(tabela.blocos_gravados())
    if len(lida.linhas) != esperadas:
        raise ErroImportacao(f"{tabela.aba}: {len(lida.linhas)} registros; esperados {esperadas}.")
    sem_valor = [linha.chave for linha in lida.linhas if not linha.origem]
    if sem_valor:
        raise ErroImportacao(f"{tabela.aba}: registros sem coluna de valor: {sem_valor[:3]}.")
    for linha in lida.linhas:
        if linha.situacao is Situacao.NAO_DISPONIVEL and linha.valor is not None:
            raise ErroImportacao(f"{linha.origem}: valor e situação incoerentes.")
