"""Gera as planilhas reduzidas de tests/fixtures/planilhas/ a partir de dados/brutos/.

Uso (dentro de backend/): uv run python tests/fixtures/gerar_planilhas.py

- anuario-reduzido.xlsx: só as abas ABAS, copiadas célula a célula com as mesclagens. Todas as
  linhas são mantidas, porque a agregação regional exige todas as UFs de cada região.
- POP2024_reduzido.ods e POP2025_reduzido.ods: só a aba BRASIL_E_UFs (sem os municípios).
- README.md: manifesto com o SHA-256 de cada arquivo, no mesmo formato de dados/README.md.

Rode de novo só quando a planilha original mudar; os testes não dependem de dados/brutos/.
"""

import hashlib
from pathlib import Path

import openpyxl
import pandas as pd

RAIZ = Path(__file__).resolve().parents[3]
ORIGEM = RAIZ / "dados" / "brutos"
DESTINO = Path(__file__).resolve().parent / "planilhas" / "brutos"
ABAS = ("T01", "T02", "T03", "T08", "T09", "T15", "T21")
IBGE = {
    "POP2024_20241230.ods": "POP2024_reduzido.ods",
    "POP2025_20260828.ods": "POP2025_reduzido.ods",
}


def copiar_abas(origem: Path, destino: Path) -> None:
    fonte = openpyxl.load_workbook(origem, data_only=True)
    novo = openpyxl.Workbook()
    novo.remove(novo.active)
    for aba in ABAS:
        ws, copia = fonte[aba], novo.create_sheet(aba)
        for linha in ws.iter_rows():
            for celula in linha:
                if celula.value is not None:
                    copia.cell(celula.row, celula.column, celula.value)
        for intervalo in ws.merged_cells.ranges:
            copia.merge_cells(str(intervalo))
    novo.save(destino)


def copiar_ibge(origem: Path, destino: Path) -> None:
    df = pd.read_excel(origem, sheet_name="BRASIL_E_UFs", engine="odf", header=None, dtype=object)
    with pd.ExcelWriter(destino, engine="odf") as escritor:
        df.to_excel(escritor, sheet_name="BRASIL_E_UFs", header=False, index=False)


def sha256(caminho: Path) -> str:
    return hashlib.sha256(caminho.read_bytes()).hexdigest()


def main() -> None:
    DESTINO.mkdir(parents=True, exist_ok=True)
    copiar_abas(ORIGEM / "anuario-2026.xlsx", DESTINO / "anuario-reduzido.xlsx")
    for original, reduzido in IBGE.items():
        copiar_ibge(ORIGEM / original, DESTINO / reduzido)
    linhas = [
        "# Planilhas reduzidas (fixtures)",
        "",
        "Geradas por `tests/fixtures/gerar_planilhas.py` a partir de `dados/brutos/`.",
        "",
        "| Arquivo | Conteúdo | SHA-256 |",
        "| --- | --- | --- |",
    ]
    for arquivo in sorted(DESTINO.iterdir()):
        linhas.append(f"| `{arquivo.name}` | reduzido | `{sha256(arquivo)}` |")
    (DESTINO.parent / "README.md").write_text("\n".join(linhas) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
