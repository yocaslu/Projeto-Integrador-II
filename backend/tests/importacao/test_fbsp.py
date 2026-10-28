"""Leitor da planilha do FBSP sobre uma aba sintética, com o mesmo padrão das tabelas reais."""

from dataclasses import replace
from decimal import Decimal

import openpyxl
import pytest

from apps.importacao.erros import ErroImportacao
from apps.importacao.leitores import fbsp
from apps.importacao.linhas import Localidade, Nivel
from apps.importacao.mapeamento import ANCORA_UFS, Bloco, Tabela
from apps.importacao.normalizacao import Situacao
from apps.registros.escolhas import Unidade

UNIDADE = {"roubo-carga": Unidade.OCORRENCIAS}
BASE = {"roubo-carga": "100 mil habitantes"}
TABELA = Tabela(
    aba="TX",
    ancora=ANCORA_UFS,
    linhas_cabecalho=3,
    blocos=(Bloco("roubo-carga", (r"^Roubo de carga$",)),),
    anos=(2024, 2025),
    localidades_esperadas=2,
)
BRASIL = Localidade(Nivel.PAIS, "BR")
ACRE = Localidade(Nivel.UF, "AC")


def criar_planilha(
    *, grupo_taxa="Taxa (2)", nota_2="(2) Taxa por 100 mil habitantes.", celulas=None, extra=None
):
    """Aba TX: título na linha 3 parecido com a âncora, cabeçalho nas linhas 5–7, dados, rodapé."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TX"
    ws["A1"] = "TABELA XX"
    ws["A3"] = "Brasil e Unidades da Federação – 2024-2025"
    ws["A5"] = "Brasil e Unidades da Federação"
    ws["B5"] = "Roubo de carga"
    ws.merge_cells("B5:F5")
    ws["B6"], ws["D6"], ws["F6"] = "Ns. Absolutos", grupo_taxa, "Variação (%)"
    ws.merge_cells("B6:C6")
    ws.merge_cells("D6:E6")
    ws.merge_cells("F6:F7")
    for coluna, ano in zip("BCDE", ("2024 (3)", 2025, 2024, 2025), strict=True):
        ws[f"{coluna}7"] = ano
    linhas = {
        9: ("Brasil", 10, 20, 1.5, 3.0000000000004, 100.0),
        10: ("Acre (4)", "-", "...", "-", "...", "..."),
    }
    for linha, valores in linhas.items():
        for coluna, valor in zip("ABCDEF", valores, strict=True):
            ws[f"{coluna}{linha}"] = valor
    ws["A12"] = "Fonte: Teste."
    ws["A13"] = "(...) Informação não disponível."
    ws["A14"] = nota_2
    ws["A15"] = "(3) Atualização das informações publicadas no Anuário."
    ws["A16"] = "(4) Nota do Acre."
    for coordenada, valor in (celulas or {}).items():
        ws[coordenada] = valor
    if extra:
        extra(ws)
    return wb


def ler(wb=None, tabela=TABELA):
    return fbsp.ler_tabela(wb or criar_planilha(), tabela, UNIDADE, BASE)


def por_chave(lida):
    return {(linha.chave.localidade, linha.chave.ano): linha for linha in lida.linhas}


def test_le_valor_taxa_situacao_e_origem():
    linhas = por_chave(ler())

    brasil = linhas[(BRASIL, 2024)]
    assert brasil.valor == Decimal(10)
    assert brasil.taxa_publicada == Decimal("1.500000000000")
    assert brasil.situacao is Situacao.INFORMADO
    assert brasil.origem == "TX!B9"
    assert brasil.chave.unidade is Unidade.OCORRENCIAS
    assert linhas[(BRASIL, 2025)].taxa_publicada == Decimal("3.000000000000")

    acre_2024, acre_2025 = linhas[(ACRE, 2024)], linhas[(ACRE, 2025)]
    assert (acre_2024.valor, acre_2024.situacao) == (Decimal(0), Situacao.INEXISTENTE)
    assert acre_2024.taxa_publicada == Decimal(0)
    assert (acre_2025.valor, acre_2025.situacao) == (None, Situacao.NAO_DISPONIVEL)
    assert acre_2025.taxa_publicada is None


def test_notas_vao_para_observacao_menos_a_da_base_da_taxa():
    linhas = por_chave(ler())

    assert linhas[(BRASIL, 2024)].observacao == "Atualização das informações publicadas no Anuário."
    assert linhas[(BRASIL, 2025)].observacao == ""
    assert linhas[(ACRE, 2025)].observacao == "Nota do Acre."
    assert "100 mil" not in linhas[(ACRE, 2024)].observacao


def test_variacao_sem_ano_no_cabecalho_e_do_ultimo_ano():
    publicados = ler().publicados

    assert {(p.grupo, p.chave.ano) for p in publicados} == {("variacao", 2025)}
    brasil = next(p for p in publicados if p.chave.localidade == BRASIL)
    assert (brasil.valor, brasil.origem) == (Decimal("100.000000000000"), "TX!F9")
    acre = next(p for p in publicados if p.chave.localidade == ACRE)
    assert acre.situacao is Situacao.NAO_DISPONIVEL


def test_base_no_rotulo_do_grupo_sem_nota():
    lida = ler(criar_planilha(grupo_taxa="Taxas por 100 mil habitantes", nota_2="(2) Outra."))
    assert len(lida.linhas) == 4


@pytest.mark.parametrize(
    ("grupo_taxa", "nota_2", "mensagem"),
    [
        ("Taxa (2)", "(2) Taxa por 100 mil veículos.", "não é a do indicador"),
        ("Taxa (5)", "(2) Taxa por 100 mil habitantes.", r"nota \(5\)"),
        ("Taxas", "(2) Taxa por 100 mil habitantes.", "não é a do indicador"),
    ],
)
def test_base_da_taxa_errada_aborta(grupo_taxa, nota_2, mensagem):
    with pytest.raises(ErroImportacao, match=mensagem):
        ler(criar_planilha(grupo_taxa=grupo_taxa, nota_2=nota_2))


def test_coluna_orfa_aborta_e_coluna_ignorada_passa():
    def outro_bloco(ws):
        ws["G5"], ws["G6"], ws["G7"] = "Outro bloco", "Ns. Absolutos", 2024

    with pytest.raises(ErroImportacao, match="TX!G5: coluna sem bloco"):
        ler(criar_planilha(extra=outro_bloco))

    lida = ler(criar_planilha(extra=outro_bloco), replace(TABELA, ignorar=(r"^Outro bloco$",)))
    assert len(lida.linhas) == 4


def test_falta_coluna_de_taxa_aborta():
    def sem_taxa(ws):
        ws.unmerge_cells("B5:F5")
        ws.unmerge_cells("D6:E6")
        ws.merge_cells("B5:C5")
        ws["F5"] = "Roubo de carga"
        for coordenada in ("D5", "D6", "D7", "E7", "D9", "E9", "D10", "E10"):
            ws[coordenada] = None

    with pytest.raises(ErroImportacao, match="falta a coluna taxa"):
        ler(criar_planilha(extra=sem_taxa))


def test_localidade_desconhecida_aborta_com_a_celula():
    with pytest.raises(ErroImportacao, match=r"TX!A10: Localidade não reconhecida"):
        ler(criar_planilha(celulas={"A10": "Atlântida"}))


def test_texto_invalido_aborta_com_a_celula():
    with pytest.raises(ErroImportacao, match=r"TX!C9: Valor não reconhecido"):
        ler(criar_planilha(celulas={"C9": "n/d"}))


def test_celula_vazia_vira_nao_disponivel_com_aviso():
    lida = ler(criar_planilha(celulas={"C9": None}))
    assert lida.vazias == ["TX!C9"]
    assert por_chave(lida)[(BRASIL, 2025)].situacao is Situacao.NAO_DISPONIVEL


def test_contagem_de_localidades_diferente_aborta():
    with pytest.raises(ErroImportacao, match="2 localidades lidas; esperadas 3"):
        ler(tabela=replace(TABELA, localidades_esperadas=3))


def test_localidade_repetida_aborta():
    with pytest.raises(ErroImportacao, match="localidade repetida"):
        ler(criar_planilha(celulas={"A10": "Brasil"}))


def test_ancora_ausente_aborta():
    with pytest.raises(ErroImportacao, match="âncora não encontrada"):
        ler(criar_planilha(celulas={"A5": "Unidades"}))


def test_tabela_de_conferencia_nao_gera_linhas():
    lida = ler(tabela=replace(TABELA, conferencia=True))
    assert lida.linhas == []
    assert {p.grupo for p in lida.publicados} == {"valor", "taxa", "variacao"}


@pytest.mark.parametrize(
    ("rotulo", "esperado"),
    [
        ("Minas Gerais (4)", ("Minas Gerais", ["4"])),
        ("Rio Grande do Norte (3) (4)", ("Rio Grande do Norte", ["3", "4"])),
        (" Brasil ", ("Brasil", [])),
        (2024, ("2024", [])),
    ],
)
def test_separar_marcas(rotulo, esperado):
    assert fbsp.separar_marcas(rotulo) == esperado
